# routes/export.py
"""Admin data exports (CSV / Excel / PDF) for Students, Attendance,
Resources, Events, Certificates, and Club Members.

Reuses existing service getters — no duplicate SQL. A single generic
export engine handles all three file formats so we never write three
near-identical implementations per dataset.
"""
import csv
import io
from datetime import date, datetime
from flask import Blueprint, Response, render_template, abort, session
from helpers.auth_helpers import admin_required
from services.user_service import get_all_users, get_user_by_id
from services.attendance_service import get_all_attendance_records
from services.resource_service import get_all_resources
from services.event_service import get_all_events
from services.certificate_service import get_all_certificates
from services.member_service import get_all_club_memberships

bp = Blueprint('export', __name__, url_prefix='/export')


def _students():
    return [u for u in get_all_users() if u['role'] == 'student']


DATASETS = {
    'students': {
        'label': 'Students',
        'columns': [('name', 'Name'), ('email', 'Email'), ('phone', 'Phone'), ('created_at', 'Joined')],
        'fetch': _students,
    },
    'attendance': {
        'label': 'Attendance',
        'columns': [('student_name', 'Student'), ('student_email', 'Email'),
                    ('event_title', 'Event'), ('club_name', 'Club'), ('scan_time', 'Marked At')],
        'fetch': get_all_attendance_records,
    },
    'resources': {
        'label': 'Resources',
        'columns': [('resource_name', 'Resource'), ('total_quantity', 'Total Qty'), ('description', 'Description')],
        'fetch': get_all_resources,
    },
    'events': {
        'label': 'Events',
        'columns': [('title', 'Title'), ('club_name', 'Club'), ('venue_name', 'Venue'),
                    ('date', 'Date'), ('start_time', 'Start'), ('end_time', 'End'),
                    ('approved_status', 'Status')],
        'fetch': lambda: get_all_events(),
    },
    'certificates': {
        'label': 'Certificates',
        'columns': [('student_name', 'Student'), ('student_email', 'Email'),
                    ('event_title', 'Event'), ('club_name', 'Club'), ('issue_date', 'Issued')],
        'fetch': get_all_certificates,
    },
    'club-members': {
        'label': 'Club Members',
        'columns': [('club_name', 'Club'), ('member_name', 'Member'), ('email', 'Email'),
                    ('position', 'Position'), ('joined_at', 'Joined')],
        'fetch': get_all_club_memberships,
    },
}


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d %H:%M') if isinstance(value, datetime) else value.strftime('%Y-%m-%d')
    return str(value)


def _export_csv(rows, columns, filename):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow([_cell(row.get(key)) for key, _ in columns])
    return Response(
        buf.getvalue(), mimetype='text/csv',
        headers={'Content-Disposition': f'attachment; filename="{filename}.csv"'}
    )


def _export_xlsx(rows, columns, filename, label):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = label[:31]

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='10B981', end_color='10B981', fill_type='solid')
    for col_idx, (_, col_label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_label)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_cell(row.get(key)))

    for col_idx, (key, col_label) in enumerate(columns, start=1):
        width = max(len(col_label), 12, *(len(_cell(r.get(key))) for r in rows)) if rows else max(len(col_label), 12)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}.xlsx"'}
    )


def _export_pdf(rows, columns, filename, label):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"CECRMS — {label} Export", styles['Title']),
        Paragraph(f"Generated {datetime.now().strftime('%d %b %Y, %I:%M %p')}", styles['Normal']),
        Spacer(1, 12),
    ]

    data = [[label for _, label in columns]]
    for row in rows:
        data.append([_cell(row.get(key)) for key, _ in columns])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#10b981')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f1f5f9')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    return Response(
        buf.read(), mimetype='application/pdf',
        headers={'Content-Disposition': f'attachment; filename="{filename}.pdf"'}
    )


@bp.route('/')
@admin_required
def index():
    user = get_user_by_id(session['user_id'])
    return render_template('export.html', user=user, active='export', datasets=DATASETS)


@bp.route('/<dataset>/<fmt>')
@admin_required
def download(dataset, fmt):
    if dataset not in DATASETS or fmt not in ('csv', 'xlsx', 'pdf'):
        abort(404)
    spec = DATASETS[dataset]
    rows = spec['fetch']()
    filename = f"cecrms_{dataset}_{date.today().isoformat()}"

    if fmt == 'csv':
        return _export_csv(rows, spec['columns'], filename)
    if fmt == 'xlsx':
        return _export_xlsx(rows, spec['columns'], filename, spec['label'])
    return _export_pdf(rows, spec['columns'], filename, spec['label'])
